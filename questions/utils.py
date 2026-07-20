import openpyxl
from .models import Question, Choice, Subject, Topic


def _resolve_subject(name, default_subject, user):
    """Cột 'Môn học' trong file: tìm theo tên (không phân biệt hoa thường), tạo mới nếu chưa có."""
    name = (name or '').strip()
    if not name:
        return default_subject
    subject = Subject.objects.filter(name__iexact=name).first()
    if not subject:
        subject = Subject.objects.create(name=name, created_by=user)
    return subject


def _resolve_topic(name, subject):
    """Cột 'Chủ đề': tìm/tạo chủ đề thuộc đúng môn học của câu hỏi."""
    name = (name or '').strip()
    if not name or not subject:
        return None
    topic, _ = Topic.objects.get_or_create(subject=subject, name=name)
    return topic


def _parse_points(value):
    try:
        pts = float(str(value).replace(',', '.'))
        return pts if pts > 0 else 1.0
    except (TypeError, ValueError):
        return 1.0


def import_questions_from_excel(file, subject, user):
    """
    Định dạng cột (hàng 1 là tiêu đề):
    A: Nội dung | B: Loại | C: Độ khó | D-H: Đáp án A-E | I: Đáp án đúng | J: Giải thích
    K: Môn học (tùy chọn — bỏ trống dùng môn đã chọn) | L: Chủ đề (tùy chọn) | M: Điểm (tùy chọn, mặc định 1)
    """
    wb = openpyxl.load_workbook(file)
    ws = wb.active
    count = 0

    for row in ws.iter_rows(min_row=2, values_only=True):
        if not row or not row[0]:
            continue
        content = str(row[0]).strip()
        if not content:
            continue

        q_type = str(row[1]).strip().lower() if row[1] else 'single'
        type_map = {
            'single': Question.TYPE_SINGLE,
            'multiple': Question.TYPE_MULTIPLE,
            'true_false': Question.TYPE_TRUE_FALSE,
            'essay': Question.TYPE_ESSAY,
            'trắc nghiệm 1 đáp án': Question.TYPE_SINGLE,
            'trắc nghiệm nhiều đáp án': Question.TYPE_MULTIPLE,
            'đúng/sai': Question.TYPE_TRUE_FALSE,
            'đúng sai': Question.TYPE_TRUE_FALSE,
            'tự luận': Question.TYPE_ESSAY,
        }
        question_type = type_map.get(q_type, Question.TYPE_SINGLE)

        difficulty_raw = str(row[2]).strip().lower() if row[2] else 'medium'
        diff_map = {'dễ': 'easy', 'easy': 'easy', 'trung bình': 'medium', 'medium': 'medium', 'khó': 'hard', 'hard': 'hard'}
        difficulty = diff_map.get(difficulty_raw, 'medium')

        choices_data = []
        correct_labels = []
        if len(row) > 3:
            labels = ['A', 'B', 'C', 'D', 'E']
            for i, label in enumerate(labels):
                col_idx = 3 + i
                if col_idx < len(row) and row[col_idx]:
                    choices_data.append((label, str(row[col_idx]).strip()))

        if len(row) > 8 and row[8]:
            correct_labels = [c.strip().upper() for c in str(row[8]).split(',')]

        explanation = str(row[9]).strip() if len(row) > 9 and row[9] else ''

        # Cột mới: K=Môn học, L=Chủ đề, M=Điểm
        row_subject = _resolve_subject(str(row[10]) if len(row) > 10 and row[10] else '', subject, user)
        topic = _resolve_topic(str(row[11]) if len(row) > 11 and row[11] else '', row_subject)
        points = _parse_points(row[12]) if len(row) > 12 else 1.0

        question = Question.objects.create(
            subject=row_subject,
            topic=topic,
            question_type=question_type,
            difficulty=difficulty,
            content=content,
            explanation=explanation,
            points=points,
            created_by=user,
        )

        for i, (label, content_c) in enumerate(choices_data):
            Choice.objects.create(
                question=question,
                label=label,
                content=content_c,
                is_correct=(label in correct_labels),
                order=i + 1,
            )
        count += 1

    return count


def import_questions_from_docx(file, subject, user):
    """
    Định dạng mỗi câu hỏi:
      Câu 1: Nội dung câu hỏi?
      A. Đáp án A
      B. Đáp án B (*đánh dấu * cho đáp án đúng hoặc ghi dòng Đáp án:)
      Đáp án: B
      Môn học: Lập trình Python   (tùy chọn — bỏ trống dùng môn đã chọn)
      Chủ đề: Kiểu dữ liệu        (tùy chọn)
      Điểm: 2                      (tùy chọn, mặc định 1)
      Độ khó: khó                  (tùy chọn: dễ/trung bình/khó)
    """
    from docx import Document
    import re

    doc = Document(file)
    count = 0
    current_question = None
    current_choices = []
    meta = {}

    diff_map = {'dễ': 'easy', 'easy': 'easy', 'trung bình': 'medium', 'medium': 'medium', 'khó': 'hard', 'hard': 'hard'}

    def save_question():
        nonlocal current_question, current_choices, meta, count
        if current_question:
            q_subject = _resolve_subject(meta.get('subject', ''), subject, user)
            topic = _resolve_topic(meta.get('topic', ''), q_subject)
            n_correct = sum(1 for _, _, ok in current_choices if ok)
            q_type = Question.TYPE_MULTIPLE if n_correct > 1 else Question.TYPE_SINGLE
            q = Question.objects.create(
                subject=q_subject,
                topic=topic,
                question_type=q_type,
                difficulty=diff_map.get(meta.get('difficulty', '').lower(), 'medium'),
                content=current_question,
                points=_parse_points(meta.get('points')) if meta.get('points') else 1.0,
                created_by=user,
            )
            for i, (label, text, is_correct) in enumerate(current_choices):
                Choice.objects.create(
                    question=q, label=label, content=text,
                    is_correct=is_correct, order=i + 1
                )
            count += 1
        current_question = None
        current_choices = []
        meta = {}

    for para in doc.paragraphs:
        text = para.text.strip()
        if not text:
            continue

        # Match question: starts with number like "1." or "Câu 1:"
        q_match = re.match(r'^(?:Câu\s+)?(\d+)[.:\)]\s+(.+)', text, re.IGNORECASE)
        if q_match:
            save_question()
            current_question = q_match.group(2).strip()
            continue

        # Answer key line: "Đáp án: A" or "Answer: A,B" (đặt TRƯỚC match A-E để không nhầm)
        ans_match = re.match(r'^(?:Đáp án|Answer)[:\s]+([A-E][A-E,\s]*)$', text, re.IGNORECASE)
        if ans_match and current_question:
            correct_labels = [c.strip().upper() for c in ans_match.group(1).split(',') if c.strip()]
            current_choices = [
                (label, content, label in correct_labels)
                for label, content, _ in current_choices
            ]
            continue

        # Metadata lines: Môn học / Chủ đề / Điểm / Độ khó
        meta_match = re.match(r'^(Môn học|Subject|Chủ đề|Topic|Điểm|Points?|Độ khó|Difficulty)[:\s]+(.+)$', text, re.IGNORECASE)
        if meta_match and current_question:
            key_raw = meta_match.group(1).lower()
            value = meta_match.group(2).strip()
            if key_raw in ('môn học', 'subject'):
                meta['subject'] = value
            elif key_raw in ('chủ đề', 'topic'):
                meta['topic'] = value
            elif key_raw in ('điểm', 'point', 'points'):
                meta['points'] = value
            else:
                meta['difficulty'] = value
            continue

        # Match choice: A. B. C. D.
        c_match = re.match(r'^([A-Ea-e])[.)\s]\s*(.+)', text)
        if c_match and current_question:
            label = c_match.group(1).upper()
            content = c_match.group(2).strip()
            is_correct = content.startswith('*') or content.endswith('*')
            content = content.strip('*').strip()
            current_choices.append((label, content, is_correct))
            continue

    save_question()
    return count
